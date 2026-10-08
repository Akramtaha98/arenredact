"""Creates Annotator_A.xlsx / Annotator_B.xlsx from items.csv (text columns, dropdowns, text-formatted cells)."""
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
items = list(csv.DictReader(open(sys.argv[1], encoding="utf-8-sig")))
for who in ("A", "B"):
    wb = Workbook(); ws = wb.active; ws.title = "annotation"
    hdr = ["item_id", "text", "no_identifier", "identifier_1", "type_1", "identifier_2", "type_2", "identifier_3", "type_3", "notes"]
    ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDE6F5")
    for it in items: ws.append([it["item_id"], it["text"]] + [None] * 8)
    for col, w in zip("ABCDEFGHIJ", (8, 70, 14, 28, 14, 28, 14, 28, 14, 30)): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        for i in (3, 5, 7): row[i].number_format = "@"     # keep digits as text (no scientific notation, keep leading zeros)
        row[1].number_format = "@"
    n = len(items) + 1
    dv1 = DataValidation(type="list", formula1='"yes"', allow_blank=True); dv1.add(f"C2:C{n}")
    dv2 = DataValidation(type="list", formula1='"PHONE,EMAIL,IBAN,NATIONAL_ID,OTHER_PII"', allow_blank=True)
    for c in "EGI": dv2.add(f"{c}2:{c}{n}")
    ws.add_data_validation(dv1); ws.add_data_validation(dv2); ws.freeze_panes = "C2"
    wb.save(f"Annotator_{who}.xlsx")
print(len(items), "items")
