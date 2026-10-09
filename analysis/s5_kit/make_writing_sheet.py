"""Create the context-writing workbook for NON-AUTHOR contributors (set S5: independent positive contexts).
Each contributor writes natural sentences with ONE slot: {ID} (a person's identifier of the stated type) or {NUM}
(a number that is NOT a personal identifier). The authors later fill slots with generated identifiers; labels
follow by construction. usage: python analysis/s5_kit/make_writing_sheet.py <outdir>"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
out = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(__file__)
TYPES = [("PHONE", "a telephone number (mobile or landline, any country in the Gulf, Iraq, Jordan or Egypt)"), ("IBAN", "a bank account number (IBAN)"),
         ("EMAIL", "an e-mail address"), ("NATIONAL_ID", "a national / civil / resident (iqama) identity number")]
LANGS = [("English", 5), ("Arabic", 5), ("Code-mixed Arabic-English", 5)]
NEG = [("English", 10), ("Arabic", 10), ("Code-mixed Arabic-English", 10)]
th = Side(style='thin', color='999999'); bd = Border(left=th, right=th, top=th, bottom=th)
for name in ("A", "B"):
    wb = Workbook(); ws = wb.active; ws.title = "Sentences"
    ws.append(["row_id", "kind", "type", "language", "what the {slot} stands for", "YOUR SENTENCE (exactly one slot)"])
    n = 0
    for t, desc in TYPES:
        for lang, k in LANGS:
            for _ in range(k):
                n += 1; ws.append([f"{name}{n:03d}", "positive", t, lang, "{ID} = " + desc, None])
    for lang, k in NEG:
        for _ in range(k):
            n += 1; ws.append([f"{name}{n:03d}", "negative", "NONE", lang, "{NUM} = a number that is NOT a personal identifier (order, ticket, shipment, price, quantity, date-like, page, flight code, ...)", None])
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.border = bd; c.alignment = Alignment(wrap_text=True, vertical='top')
        row[5].fill = PatternFill('solid', fgColor='F1F8E9')
    for col, w in zip("ABCDEF", [8, 10, 14, 24, 60, 80]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    g = wb.create_sheet("Instructions")
    for l in ["Writing natural contexts for a redaction test (90 sentences, about 45 minutes)",
      "You are NOT an author of the paper and you do not need to know how the detector works. Please do not look at any detector, list or earlier test set.",
      "For each row write ONE natural sentence a real person might write or say (chat message, e-mail line, form field, note, call transcript, table cell) in the stated language.",
      "Put the slot exactly once, as {ID} for positive rows or {NUM} for negative rows. The authors will insert a generated identifier or number there. Example: 'Please call me on {ID} after lunch.'",
      "Positive rows: the slot must be where the identifier of the stated type would really appear, and the sentence should make clear, in your own words, what it is (as you would naturally say it).",
      "Negative rows: write a sentence in which a number of similar length appears but it is clearly NOT a personal identifier. Include some that use identity-like words (for example 'identity documents', 'the card reader') without an identity number.",
      "Vary style, length and vocabulary. Do not reuse one template. Do not put real people's data in. Do not include any other numbers or e-mail addresses in the sentence.",
      "Arabic rows should be natural Arabic (Gulf, Levantine, Egyptian or Iraqi usage is welcome); code-mixed rows should mix Arabic and English the way you or your colleagues do.",
      "Return the workbook unchanged except for the last column."]: g.append([l])
    g.column_dimensions['A'].width = 140
    for r in g.iter_rows(): r[0].alignment = Alignment(wrap_text=True)
    g['A1'].font = Font(bold=True, size=13)
    wb.save(os.path.join(out, f"S5_Writing_Sheet_{name}.xlsx"))
print("ok", out)
