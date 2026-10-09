"""Build the blinded human review sheet for the leakage scorer from REAL attacked outputs of S3.
120 items = 2 systems (ours, Presidio+S1+reg; hidden) x scorer verdict (leaked / not leaked; hidden) x 30,
drawn at random (seed 20261012) from identifiers that were altered by one of 10 attacks and protected on clean text.
Reviewers see: identifier type, the ORIGINAL identifier and the redacted output. They do not see system, attack or scorer verdict.
usage: python analysis/make_leakage_sheet.py -> analysis/leakage_review/Leakage_Review_Sheet.xlsx (+ authors-only key CSV)"""
import csv, os, random, sys
sys.path.insert(0, os.path.dirname(__file__))
import common as C
L = C.L
import run_attack_audit as RA
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
CFGS = ["HGL 100%", "INV (dev chars)", "BIDI controls", "Fullwidth/math forms", "Arabic-Indic digits", "SEP (dev separators)", "SEP (held-out separators)", "CONF (folded set)", "CONF (held-out set)", "COMBO-2"]
sets = RA.load_sets(); cases = sets["stratified_S3"]
sysd = C.make_systems(False); sysd = {"A": sysd["ours_full"], "B": sysd["presidio_norm_regional"]}
pool = {(s, v): [] for s in sysd for v in (True, False)}
cfgs = [c for c in RA.FIXED if c[0] in CFGS]
for cfg in cfgs:
    for i, (t, g) in enumerate(cases):
        a = RA.attack_sentence(t, sorted(g), cfg, (cfg[0], "stratified_S3", i))[0]
        for sname, sysobj in sysd.items():
            _, _, red0 = sysobj.run(t); _, _, red = sysobj.run(a.text)
            for j, (s, e, ty) in enumerate(sorted(g)):
                if not a.altered[j]: continue
                if L.leaks(t[s:e], red0, context=t[:s] + t[e:]): continue
                ctx = a.text[:a.gold[j][0]] + a.text[a.gold[j][1]:]
                pool[(sname, L.leaks(t[s:e], red, context=ctx))].append(dict(system=sname, attack=cfg[0], i=i, j=j, type=ty, ident=t[s:e], output=red,
                    partial=L.leaks(t[s:e], red, context=ctx), full=L.leaks(t[s:e], red, k=10**6, context=ctx)))
rng = random.Random(20261012); items = []
for k, v in pool.items():
    print(k, len(v)); items += rng.sample(v, min(30, len(v)))
rng.shuffle(items)
os.makedirs(os.path.join(C.ROOT, "analysis/leakage_review/_private"), exist_ok=True)
wb = Workbook(); ws = wb.active; ws.title = "Review"
ws.append(["item_id", "identifier type", "ORIGINAL identifier", "Redacted output (what a recipient would see)", "Q1: Could you reconstruct the COMPLETE identifier from the output? (yes/no)", "Q2: Does the output reveal enough of the identifier that you would count it as a privacy leak? (yes/no)", "Notes (optional)"])
key = []
for n, it in enumerate(items, 1):
    iid = f"L{n:03d}"; ws.append([iid, it["type"], it["ident"], it["output"], None, None, None])
    key.append([iid, it["system"], it["attack"], it["type"], int(it["partial"]), int(it["full"])])
th = Side(style='thin', color='999999'); bd = Border(left=th, right=th, top=th, bottom=th)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.alignment = Alignment(wrap_text=True, vertical='center'); c.fill = PatternFill('solid', fgColor='1F3864' if c.column <= 4 else '2E7D32')
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = bd
        if c.column >= 5: c.fill = PatternFill('solid', fgColor='F1F8E9')
for col, w in zip("ABCDEFG", [8, 12, 34, 70, 22, 26, 30]): ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 75; ws.freeze_panes = "D2"
dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"E2:F{len(items)+1}")
g = wb.create_sheet("Instructions")
for l in ["Leakage review (%d items)" % len(items),
 "You are shown a sentence that a redaction tool has already processed, and the ORIGINAL identifier it was meant to hide. Do not look anything up and do not use any tool.",
 "Q1: Using only the redacted output, could you reconstruct the COMPLETE original identifier (every character, in order)? yes / no.",
 "Q2: Would you, as a privacy officer, count this output as a leak of the identifier (enough of it is readable to link, guess or recover it)? yes / no. A few scattered characters that are common in the surrounding text do not count.",
 "Note: characters that look alike (for example Arabic-Indic and Western digits, or Cyrillic letters that look Latin) count as the same character for reading purposes.",
 "Work alone. You are not told which tool produced an output or how it was altered. Do not discuss items with the authors before returning the sheet.",
 "You must not be an author of the paper. About 40 minutes."]: g.append([l])
g.column_dimensions['A'].width = 150
for r in g.iter_rows(): r[0].alignment = Alignment(wrap_text=True)
g['A1'].font = Font(bold=True, size=13)
wb.save(os.path.join(C.ROOT, "analysis/leakage_review/Leakage_Review_Sheet.xlsx"))
with open(os.path.join(C.ROOT, "analysis/leakage_review/_private/AUTHORS_ONLY_key.csv"), "w", newline="", encoding="utf8") as f:
    w = csv.writer(f); w.writerow(["item_id", "system(A=ours,B=Presidio+S1+reg)", "attack", "type", "scorer_partial_k6", "scorer_full"]); w.writerows(key)
print("items", len(items))
